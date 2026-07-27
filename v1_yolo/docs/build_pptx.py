from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x05, 0x0B, 0x1F)
NAVY_PANEL = RGBColor(0x0B, 0x14, 0x33)
BLUE = RGBColor(0x2A, 0x6F, 0xDB)
BLUE_LIGHT = RGBColor(0x6F, 0xA8, 0xF5)
WHITE = RGBColor(0xE8, 0xEC, 0xF5)
MUTED = RGBColor(0x7C, 0x90, 0xB8)

LOGO = "/home/rakine/osc/v1_yolo/docs/logo_icon.png"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False
    # renvoyer au fond
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return slide


def add_text(slide, left, top, width, height, text, size, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Segoe UI", spacing=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if spacing:
        p.line_spacing = spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return box


def add_lines(slide, left, top, width, height, lines, size, color=WHITE, bold=False,
              align=PP_ALIGN.LEFT, space_after=10, font="Segoe UI"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return box


def accent_line(slide, left, top, width, color=BLUE, height=Pt(3)):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    line.shadow.inherit = False
    return line


def footer(slide, page_label):
    add_text(slide, Inches(11.7), Inches(7.05), Inches(1.4), Inches(0.35),
              page_label, 10, MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, Inches(0.5), Inches(7.05), Inches(4), Inches(0.35),
              "SENTINEL X", 10, MUTED)


def kicker_title(slide, kicker, title, top=Inches(0.55)):
    add_text(slide, Inches(0.7), top, Inches(8), Inches(0.4), kicker.upper(), 13, BLUE_LIGHT, bold=True)
    add_text(slide, Inches(0.7), top + Inches(0.4), Inches(11), Inches(0.9), title, 34, WHITE, bold=True)
    accent_line(slide, Inches(0.72), top + Inches(1.25), Inches(0.9))


# ---------- Slide 1 : Titre ----------
s = add_slide()
s.shapes.add_picture(LOGO, Inches(4.77), Inches(1.1), height=Inches(2.0))
add_text(s, Inches(0), Inches(3.85), prs.slide_width, Inches(0.9),
          "SENTINEL X", 54, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(4.55), prs.slide_width, Inches(0.5),
          "VOIR AVANT — AGIR À TEMPS", 16, BLUE_LIGHT, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(6.4), prs.slide_width, Inches(0.5),
          "Tech Impact — Orange Summer Challenge 2026", 14, MUTED, align=PP_ALIGN.CENTER)

# ---------- Slide 2 : Le problème ----------
s = add_slide()
kicker_title(s, "Constat", "La vidéosurveillance ne comprend rien à ce qu'elle filme")
add_lines(s, Inches(0.7), Inches(2.6), Inches(9.5), Inches(3.5), [
    "Une caméra enregistre. Un détecteur de mouvement réagit à un déplacement,",
    "jamais à une situation.",
    "",
    "Résultat : on découvre toujours après coup — une fois le vol,",
    "le cambriolage ou l'agression déjà commis.",
], 22, WHITE, space_after=6)
footer(s, "2")

# ---------- Slide 3 : La mission ----------
s = add_slide()
kicker_title(s, "Mission", "Faire passer la caméra d'un témoin à une vigie")
principles = [
    ("Comportement, pas identité", "Observer la posture et les gestes, pas reconnaître un visage."),
    ("Signaux avant l'acte", "Repérer ce qui précède une situation à risque."),
    ("Un humain décide", "Le système alerte, il ne décide jamais seul."),
    ("Réponse selon la gravité", "Notification simple, ou appel vocal si c'est critique."),
]
x0 = Inches(0.7)
y0 = Inches(2.5)
card_w = Inches(2.85)
gap = Inches(0.22)
for i, (title, sub) in enumerate(principles):
    x = x0 + i * (card_w + gap)
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y0, card_w, Inches(3.4))
    card.fill.solid()
    card.fill.fore_color.rgb = NAVY_PANEL
    card.line.color.rgb = RGBColor(0x1B, 0x2A, 0x55)
    card.line.width = Pt(1)
    card.shadow.inherit = False
    card.adjustments[0] = 0.06
    add_text(s, x + Inches(0.18), y0 + Inches(0.3), card_w - Inches(0.36), Inches(0.4),
              f"0{i+1}", 20, BLUE_LIGHT, bold=True)
    add_text(s, x + Inches(0.18), y0 + Inches(0.85), card_w - Inches(0.36), Inches(1.1),
              title, 17, WHITE, bold=True)
    add_text(s, x + Inches(0.18), y0 + Inches(2.0), card_w - Inches(0.36), Inches(1.3),
              sub, 12.5, MUTED)
footer(s, "3")

# ---------- Slide 4 : Comment ça marche ----------
s = add_slide()
kicker_title(s, "Fonctionnement", "Une chaîne d'analyse, pas un simple détecteur")
steps = [
    ("Caméra", "IP ou téléphone, réseau existant"),
    ("Chaîne d'analyse", "13 modèles en parallèle + fiche par personne"),
    ("Alerte graduée", "Notification, ou appel vocal si critique"),
]
box_w = Inches(3.3)
box_y = Inches(3.1)
gap2 = Inches(1.0)
total_w = box_w * 3 + gap2 * 2
x0 = (prs.slide_width - total_w) / 2
for i, (title, sub) in enumerate(steps):
    x = x0 + i * (box_w + gap2)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, box_y, box_w, Inches(1.6))
    box.fill.solid()
    box.fill.fore_color.rgb = NAVY_PANEL
    box.line.color.rgb = BLUE
    box.line.width = Pt(1.25)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(6)
    r2 = p2.add_run()
    r2.text = sub
    r2.font.size = Pt(12)
    r2.font.color.rgb = MUTED
    if i < 2:
        arrow_x = x + box_w + Inches(0.08)
        arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, arrow_x, box_y + Inches(0.55), Inches(0.85), Inches(0.5))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = BLUE
        arrow.line.fill.background()
        arrow.shadow.inherit = False
add_text(s, Inches(0.7), Inches(5.35), Inches(11.9), Inches(0.9),
          "Le système ne décrit que des faits observés — jamais une supposition présentée comme un fait.",
          15, BLUE_LIGHT, align=PP_ALIGN.CENTER)
footer(s, "4")

# ---------- Slide 5 : Les rôles donnés à l'IA ----------
s = add_slide()
kicker_title(s, "La chaîne d'IA", "Chaque modèle a un rôle précis, pas une boîte noire")
roles = [
    ("La Sentinelle", "Armes et objets dangereux"),
    ("Le Guetteur", "Feu et fumée"),
    ("L'Arbitre", "Bagarre, comportement violent"),
    ("Le Physionomiste", "Reconnaît un visage déjà croisé"),
    ("Le Profileur", "Fiche par personne : genre, expression, tenue"),
    ("L'Inspecteur", "Casque, gilet, équipement de sécurité"),
]
rcol_w = Inches(3.85)
rcard_h = Inches(1.15)
rx0 = Inches(0.7)
ry0 = Inches(2.35)
rgap_x = Inches(0.25)
rgap_y = Inches(0.2)
for i, (title, sub) in enumerate(roles):
    col = i % 3
    row = i // 3
    x = rx0 + col * (rcol_w + rgap_x)
    y = ry0 + row * (rcard_h + rgap_y)
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, rcol_w, rcard_h)
    card.fill.solid()
    card.fill.fore_color.rgb = NAVY_PANEL
    card.line.color.rgb = RGBColor(0x1B, 0x2A, 0x55)
    card.line.width = Pt(1)
    card.shadow.inherit = False
    card.adjustments[0] = 0.1
    add_text(s, x + Inches(0.2), y + Inches(0.12), rcol_w - Inches(0.4), Inches(0.4), title, 15, BLUE_LIGHT, bold=True)
    add_text(s, x + Inches(0.2), y + Inches(0.55), rcol_w - Inches(0.4), Inches(0.55), sub, 11.5, WHITE)

note_y = ry0 + 2 * (rcard_h + rgap_y) + Inches(0.15)
note = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, rx0, note_y, Inches(11.9), Inches(0.9))
note.fill.solid()
note.fill.fore_color.rgb = NAVY_PANEL
note.line.color.rgb = BLUE
note.line.width = Pt(1)
note.shadow.inherit = False
add_lines(s, rx0 + Inches(0.25), note_y + Inches(0.1), Inches(11.4), Inches(0.75), [
    "Chaque caméra journalise ses détections et ses narrations en JSON.",
    "Un visage jamais vu devient « Inconnu 1, 2, 3... » — nommable à tout moment, sans jamais être perdu.",
], 12.5, MUTED, space_after=2)
footer(s, "5")

# ---------- Slide 6 : Ce qui est fait ----------
s = add_slide()
kicker_title(s, "Avancement", "Ce qui est fait aujourd'hui")
done = [
    "Multi-caméra, multi-bâtiment, accès par code",
    "Détection comportementale : menace, feu, bagarre, chute",
    "Reconnaissance faciale + autorisations par lieu",
    "Genre, expression, vêtements, équipement de sécurité",
    "Présence prolongée détectée (rôdeur potentiel)",
    "Assistant conversationnel — réponses basées sur les données réelles",
]
col_w = Inches(5.7)
for i, item in enumerate(done):
    col = 0 if i < 3 else 1
    row = i % 3
    x = Inches(0.7) + col * (col_w + Inches(0.3))
    y = Inches(2.5) + row * Inches(1.15)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y + Inches(0.12), Inches(0.14), Inches(0.14))
    dot.fill.solid()
    dot.fill.fore_color.rgb = BLUE
    dot.line.fill.background()
    dot.shadow.inherit = False
    add_text(s, x + Inches(0.3), y, col_w - Inches(0.3), Inches(1.0), item, 16, WHITE)
footer(s, "6")

# ---------- Slide 7 : Prochaines étapes ----------
s = add_slide()
kicker_title(s, "Prochaines étapes", "Ce qu'on teste et ajuste maintenant")
todo = [
    ("Tests en conditions réelles", "Plusieurs personnes, plusieurs lieux, sur la durée."),
    ("Réglage des seuils", "Confiance des modèles, durée avant l'alerte rôdeur, faux positifs."),
    ("Modèle cagoule / cache-nez", "En cours d'entraînement, pas encore intégré à la chaîne."),
]
y0 = Inches(2.6)
for i, (title, sub) in enumerate(todo):
    y = y0 + i * Inches(1.35)
    num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), y, Inches(0.55), Inches(0.55))
    num.fill.solid()
    num.fill.fore_color.rgb = NAVY_PANEL
    num.line.color.rgb = BLUE
    num.line.width = Pt(1.5)
    num.shadow.inherit = False
    tf = num.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(i + 1)
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = BLUE_LIGHT
    add_text(s, Inches(1.5), y - Inches(0.05), Inches(10), Inches(0.5), title, 19, WHITE, bold=True)
    add_text(s, Inches(1.5), y + Inches(0.5), Inches(10), Inches(0.5), sub, 13.5, MUTED)
footer(s, "7")

# ---------- Slide 8 : Clôture ----------
s = add_slide()
s.shapes.add_picture(LOGO, Inches(5.15), Inches(1.6), height=Inches(1.6))
add_text(s, Inches(0), Inches(3.55), prs.slide_width, Inches(0.8),
          "SENTINEL X", 40, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(4.15), prs.slide_width, Inches(0.5),
          "VOIR AVANT — AGIR À TEMPS", 14, BLUE_LIGHT, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(5.3), prs.slide_width, Inches(0.5),
          "Merci", 20, MUTED, align=PP_ALIGN.CENTER)

prs.save("/home/rakine/osc/v1_yolo/docs/SENTINELX_Presentation.pptx")
print("OK")
